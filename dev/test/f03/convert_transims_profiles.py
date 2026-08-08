"""Regenerate transims/departure_profile.csv from the source workbook.

Source: OpenDTA_ASU/consensus_datasets/sample_departure_time_profiles.xlsx,
sheet "5.References", rows 44-51: seven diurnal profiles (HBW_SOV_PA/AP,
HBO_SOV_AP/PA, NHB_SOV_AP/PA, Trucks), 96 x 15-min bins, 5h00 wrapping past
midnight to 4h45. Cell C54 credits the FHWA Alexandria TRANSIMS Data Set.

These are the full originals of the truncated (93-bin) profiles embedded in
data/Chicago_Sketch/settings.yml — the repo copies drop the 4:15/4:30/4:45
bins, which is exactly the Sigma gap flagged in
dev/test/gold/reports/chicago/chicago_profile_audit.csv.

Deterministic, RNG-free. Bins before 5 AM are emitted on the monotone >24h
clock (24:00-28:45) per the gold v2 convention. Raw weights are written to
`weight` so the engine's own G5 normalization is exercised (Trucks sums to
0.99989 -> REPAIRED tier with recorded factor).
"""
import csv
import openpyxl

SRC = r"..\..\..\..\consensus_datasets\sample_departure_time_profiles.xlsx"
DST = r"transims\departure_profile.csv"

HDR_ROW, FIRST_ROW, LAST_ROW = 44, 45, 51
PERIOD_ID = 1  # single all-day demand period in the fixture's settings.yml

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["5.References"]

cols = []
for c in range(3, 120):
    v = ws.cell(row=HDR_ROW, column=c).value
    if v is None or "h" not in str(v):
        continue
    cols.append((c, str(v)))


def to_monotone_sec(label):
    h, m = (int(x) for x in label.split("h"))
    if h < 5:  # horizon starts 05:00; earlier hours belong to the next day
        h += 24
    return h * 3600 + m * 60


with open(DST, "w", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["profile_id", "period_id", "departure_time", "bin_width_sec",
                "weight", "weight_raw", "profile_source"])
    for r in range(FIRST_ROW, LAST_ROW + 1):
        pid = ws.cell(row=r, column=3).value
        bins = sorted((to_monotone_sec(label), float(ws.cell(row=r, column=c).value))
                      for c, label in cols)
        for sec, wt in bins:
            hh, rem = divmod(sec, 3600)
            t = f"{hh:02d}:{rem // 60:02d}:00"
            w.writerow([pid, PERIOD_ID, t, 900, repr(wt), repr(wt),
                        "FHWA_Alexandria_TRANSIMS_sheet5_references"])

print(f"wrote {DST}")
